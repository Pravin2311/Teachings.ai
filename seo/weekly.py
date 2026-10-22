"""
Weekly SEO learning loop for teachings.ai: fetch the latest complete Search
Console, GA4 and AdSense data (read-only), save a dated snapshot, evaluate
logged experiments, and write a dated report with a prioritized plan.

Usage (from the repo root, with the seo venv):
  seo\\.venv\\Scripts\\python seo\\weekly.py --auth          # one-time browser consent
  seo\\.venv\\Scripts\\python seo\\weekly.py --check         # test each API against real data
  seo\\.venv\\Scripts\\python seo\\weekly.py                 # full weekly run
  seo\\.venv\\Scripts\\python seo\\weekly.py --from-export teachings.ai-Performance-on-Search-2026-10-18.xlsx
                                                         # no API: build the report from a
                                                         # manual Search Console export

Nothing here writes to Google, the site, or ads. Snapshots and reports go to
seo/data/ and seo/reports/ (git-ignored: they contain revenue figures).
"""
import argparse
import datetime as dt
import glob
import json
import os
import re
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA_DIR = os.path.join(HERE, "data")
REPORT_DIR = os.path.join(HERE, "reports")
EXPERIMENTS = os.path.join(HERE, "experiments.json")
sys.path.insert(0, HERE)

import analysis  # noqa: E402


def load_config():
    with open(os.path.join(HERE, "config.json"), encoding="utf-8") as f:
        return json.load(f)


def windows(end):
    d = dt.timedelta
    return {
        "cur7": (end - d(6), end), "prev7": (end - d(13), end - d(7)),
        "cur28": (end - d(27), end), "prev28": (end - d(55), end - d(28)),
        # same weekdays one year earlier (364 days) -- only meaningful once the site is a year old
        "yoy28": (end - d(27 + 364), end - d(364)),
    }


def sitemap_urls():
    with open(os.path.join(ROOT, "sitemap.xml"), encoding="utf-8") as f:
        return re.findall(r"<loc>(.*?)</loc>", f.read())


# ---------------------------------------------------------------- API fetch

def pick_gsc_site(gsc, cfg):
    if cfg.get("gsc_site"):
        return cfg["gsc_site"]
    sites = {s["siteUrl"]: s.get("permissionLevel") for s in gsc.sites()}
    for cand in ("sc-domain:teachings.ai", "https://www.teachings.ai/", "https://teachings.ai/"):
        if cand in sites and sites[cand] != "siteUnverifiedUser":
            return cand
    raise RuntimeError(f"No verified Search Console property for teachings.ai among: {list(sites)}")


def fetch_gsc(gapi, creds, cfg, win, inspect_n):
    gsc = gapi.SearchConsole(creds)
    site = pick_gsc_site(gsc, cfg)
    out = {"site": site, "daily": [], "pages": {}, "queries": {}, "page_query": {},
           "countries": {}, "devices": {}}
    # 16 months is the most Search Console keeps
    out["daily"] = gsc.query(site, win["cur28"][1] - dt.timedelta(days=486), win["cur28"][1], ["date"])
    for key in ("cur7", "prev7", "cur28", "prev28"):
        s, e = win[key]
        out["pages"][key] = gsc.query(site, s, e, ["page"])
        out["queries"][key] = gsc.query(site, s, e, ["query"])
    for key in ("cur28", "prev28"):
        s, e = win[key]
        out["page_query"][key] = gsc.query(site, s, e, ["page", "query"])
        out["countries"][key] = gsc.query(site, s, e, ["country"])
        out["devices"][key] = gsc.query(site, s, e, ["device"])
    try:
        out["sitemaps"] = gsc.sitemaps(site)
    except gapi.HttpError as e:
        out["sitemaps_error"] = gapi.explain_http_error(e)
    # URL inspection: top pages by impressions plus sitemap URLs with no impressions at all
    if inspect_n:
        seen = {r["page"] for r in out["pages"]["cur28"]}
        top = [r["page"] for r in sorted(out["pages"]["cur28"], key=lambda r: -r["impressions"])]
        unseen = [u for u in sitemap_urls() if u not in seen]
        watch = [cfg["site_origin"] + u for u in cfg.get("watch_paths", [])]
        targets = watch + [u for u in top if u not in watch][: inspect_n // 2] + unseen[: inspect_n - inspect_n // 2]
        out["inspections"] = []
        for u in targets:
            try:
                out["inspections"].append(gsc.inspect(site, u))
            except gapi.HttpError as e:
                out["inspections"].append({"url": u, "error": gapi.explain_http_error(e)})
                break
            time.sleep(0.2)
    return out


def fetch_ga4(gapi, creds, cfg, win):
    ga = gapi.Analytics(creds)
    prop = cfg.get("ga4_property_id") or ga.find_property(cfg["ga4_measurement_id"])
    if not prop:
        raise RuntimeError(f"No GA4 property visible to this account has a web stream "
                           f"with measurement id {cfg['ga4_measurement_id']}")
    m_core = ["totalUsers", "sessions", "engagedSessions", "screenPageViews"]
    out = {"property": prop, "totals": {}, "landing_organic": {}}
    s0, e0 = win["prev28"][0], win["cur28"][1]
    out["daily_channels"] = ga.report(prop, s0, e0, ["date", "sessionDefaultChannelGroup"], m_core)
    for key in ("cur7", "prev7", "cur28", "prev28"):
        s, e = win[key]
        # no date dimension -> users are de-duplicated across the window
        out["totals"][key] = ga.report(prop, s, e, ["sessionDefaultChannelGroup"], m_core)
    for key in ("cur28", "prev28"):
        s, e = win[key]
        out["landing_organic"][key] = ga.report(
            prop, s, e, ["landingPage"], m_core + ["engagementRate", "userEngagementDuration"],
            organic_only=True, order_by="sessions", limit=500)
    s, e = win["cur28"]
    out["countries_organic"] = ga.report(prop, s, e, ["country"], m_core, organic_only=True, limit=50,
                                         order_by="totalUsers")
    out["devices_organic"] = ga.report(prop, s, e, ["deviceCategory"], m_core, organic_only=True)
    # page-level ad revenue only works if AdSense is linked to this GA4 property
    try:
        out["ad_pages"] = ga.report(prop, s, e, ["pagePath"],
                                    ["screenPageViews", "publisherAdImpressions", "totalAdRevenue"],
                                    order_by="screenPageViews", limit=500)
    except gapi.HttpError as ex:
        out["ad_pages_error"] = gapi.explain_http_error(ex)
    return out


def fetch_adsense(gapi, creds, cfg, win):
    ads = gapi.AdSense(creds)
    acct = cfg.get("adsense_account")
    names = [a["name"] for a in ads.accounts()]
    if acct not in names:
        if not names:
            raise RuntimeError("Signed-in account has no AdSense account")
        acct = names[0]
    metrics = ["ESTIMATED_EARNINGS", "PAGE_VIEWS", "PAGE_VIEWS_RPM", "IMPRESSIONS", "CLICKS"]
    s0, e0 = win["prev28"][0], win["cur28"][1]
    out = {"account": acct}
    try:
        out["daily"] = ads.report(acct, s0, e0, ["DATE", "DOMAIN_NAME"], metrics)
    except gapi.HttpError:
        out["daily"] = ads.report(acct, s0, e0, ["DATE"], metrics)
    s, e = win["cur28"]
    out["countries"] = ads.report(acct, s, e, ["COUNTRY_NAME"], metrics)
    return out


def fetch_all(cfg, end, inspect_n, only=None):
    import gapi
    creds = gapi.get_credentials(interactive=False)
    win = windows(end)
    snap = {"meta": {"source": "api", "end": end.isoformat(),
                     "fetched_at": dt.datetime.now().isoformat(timespec="seconds")},
            "windows": {k: [a.isoformat(), b.isoformat()] for k, (a, b) in win.items()},
            "errors": {}}
    jobs = {"gsc": lambda: fetch_gsc(gapi, creds, cfg, win, inspect_n),
            "ga4": lambda: fetch_ga4(gapi, creds, cfg, win),
            "adsense": lambda: fetch_adsense(gapi, creds, cfg, win)}
    for name, job in jobs.items():
        if only and name not in only:
            continue
        try:
            snap[name] = job()
        except gapi.HttpError as e:
            snap[name] = None
            snap["errors"][name] = gapi.explain_http_error(e)
        except Exception as e:  # report every source's failure honestly, keep the others
            snap[name] = None
            snap["errors"][name] = f"{type(e).__name__}: {e}"
    return snap


# ---------------------------------------------------------------- export fallback

def snapshot_from_export(path):
    """Search Console UI export (.xlsx). Daily totals only support period comparisons;
    page/query/country/device tables are aggregates over the whole export range."""
    import openpyxl
    import warnings
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(path, read_only=True)

    def table(sheet, key):
        rows = list(wb[sheet].iter_rows(values_only=True))[1:]
        return [{key: r[0], "clicks": r[1] or 0, "impressions": r[2] or 0,
                 "ctr": r[3] or 0, "position": r[4] or 0} for r in rows if r[0]]

    daily = [{"date": r["date"], **{k: r[k] for k in ("clicks", "impressions", "ctr", "position")}}
             for r in table("Chart", "date")]
    end = dt.date.fromisoformat(daily[-1]["date"])
    win = windows(end)
    filters = dict(list(wb["Filters"].iter_rows(values_only=True))[1:])
    return {
        "meta": {"source": "export", "export_file": os.path.basename(path), "end": end.isoformat(),
                 "export_range": filters.get("Date"), "fetched_at": dt.datetime.now().isoformat(timespec="seconds")},
        "windows": {k: [a.isoformat(), b.isoformat()] for k, (a, b) in win.items()},
        "gsc": {"site": "export", "daily": daily,
                "pages": {"export": table("Pages", "page")},
                "queries": {"export": table("Queries", "query")},
                "countries": {"export": table("Countries", "country")},
                "devices": {"export": table("Devices", "device")}},
        "ga4": None, "adsense": None,
        "errors": {"ga4": "Not fetched: export mode (API not connected yet)",
                   "adsense": "Not fetched: export mode (API not connected yet)"},
    }


# ---------------------------------------------------------------- main

def check(cfg):
    """Smoke-test every integration against real data and say exactly what failed."""
    end = dt.date.today() - dt.timedelta(days=cfg["gsc_lag_days"])
    try:
        snap = fetch_all(cfg, end, inspect_n=1)
    except Exception as e:
        print(f"[FAIL] {e}")
        return False
    ok = True
    g = snap.get("gsc")
    if g:
        c = sum(r["clicks"] for r in g["pages"]["cur28"])
        print(f"[ ok ] Search Console {g['site']}: {len(g['daily'])} days, {c:.0f} clicks in last 28d")
    a = snap.get("ga4")
    if a:
        u = sum(r["totalUsers"] for r in a["totals"]["cur28"])
        print(f"[ ok ] GA4 property {a['property']}: {u:.0f} users in last 28d")
    s = snap.get("adsense")
    if s:
        rev = sum(r["ESTIMATED_EARNINGS"] for r in s["daily"])
        cur = s["daily"][0]["currency"] if s["daily"] else ""
        print(f"[ ok ] AdSense {s['account']}: {len(s['daily'])} daily rows, {cur} {rev:.2f} over 56d")
    for name, err in snap["errors"].items():
        ok = False
        print(f"[FAIL] {name}: {err}")
    return ok


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--auth", action="store_true", help="run the one-time OAuth consent in a browser")
    ap.add_argument("--check", action="store_true", help="test each API against real data")
    ap.add_argument("--from-export", metavar="XLSX", help="build the report from a Search Console export")
    ap.add_argument("--end", help="last day to include (YYYY-MM-DD); default today minus gsc_lag_days")
    ap.add_argument("--inspect", type=int, help="URL inspections this run (quota 2000/day)")
    ap.add_argument("--no-write-experiments", action="store_true", help="don't update experiments.json")
    args = ap.parse_args()
    cfg = load_config()

    if args.auth:
        import gapi
        try:
            gapi.get_credentials(interactive=True)
        except gapi.SetupError as e:
            print(f"Setup needed: {e}", file=sys.stderr)
            return 2
        print(f"Token saved under {gapi.CREDS_DIR}. Now run: seo\\.venv\\Scripts\\python seo\\weekly.py --check")
        return 0
    if args.check:
        return 0 if check(cfg) else 1

    if args.from_export:
        snap = snapshot_from_export(args.from_export)
    else:
        end = (dt.date.fromisoformat(args.end) if args.end
               else dt.date.today() - dt.timedelta(days=cfg["gsc_lag_days"]))
        inspect_n = cfg["url_inspections_per_run"] if args.inspect is None else args.inspect
        try:
            snap = fetch_all(cfg, end, inspect_n)
        except Exception as e:  # missing token etc. -- nothing was fetched
            print(f"Could not fetch data: {e}", file=sys.stderr)
            return 2

    run_date = dt.date.today().isoformat()
    snap_dir = os.path.join(DATA_DIR, run_date)
    os.makedirs(snap_dir, exist_ok=True)
    with open(os.path.join(snap_dir, f"snapshot-{snap['meta']['source']}.json"), "w", encoding="utf-8") as f:
        json.dump(snap, f, indent=1, default=str)

    with open(EXPERIMENTS, encoding="utf-8") as f:
        experiments = json.load(f)
    gsc_client = None
    if snap["meta"]["source"] == "api" and snap.get("gsc"):
        import gapi
        gsc_client = (gapi.SearchConsole(gapi.get_credentials()), snap["gsc"]["site"])
    exp_results = analysis.evaluate_experiments(experiments, snap, gsc_client)
    if not args.no_write_experiments and gsc_client:
        with open(EXPERIMENTS, "w", encoding="utf-8") as f:
            json.dump(experiments, f, indent=2)
            f.write("\n")

    previous = sorted(glob.glob(os.path.join(REPORT_DIR, "*-plan.json")))
    prev_plan = None
    if previous:
        with open(previous[-1], encoding="utf-8") as f:
            prev_plan = json.load(f)

    result = analysis.analyze(snap, cfg, sitemap_urls(), experiments)
    md, plan = analysis.render_report(snap, cfg, result, exp_results, prev_plan, run_date)
    os.makedirs(REPORT_DIR, exist_ok=True)
    suffix = "" if snap["meta"]["source"] == "api" else "-export"
    report_path = os.path.join(REPORT_DIR, f"{run_date}{suffix}.md")
    with open(report_path, "w", encoding="utf-8") as f:
        f.write(md)
    with open(os.path.join(REPORT_DIR, f"{run_date}{suffix}-plan.json"), "w", encoding="utf-8") as f:
        json.dump(plan, f, indent=2)
    print(f"Report: {report_path}")
    for name, err in snap["errors"].items():
        print(f"[missing] {name}: {err}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
